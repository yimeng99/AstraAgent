"""把上传的文档（Word / Excel / PDF / 纯文本）抽成模型能读的纯文本。

**为什么必须在服务端抽**：

* 模型只吃文本 + 图片/视频。``.docx`` / ``.xlsx`` 是 zip 容器、``.pdf`` 是
  二进制流，把字节直接塞给它等于让它猜；
* 让模型自己在沙箱里解析也不可靠 —— 沙箱镜像 ``python:3.11-slim`` 里没有
  这些库，容器内也装不了。

所以流程是：**上传时只落盘 → 进对话前在这里抽成文本**（见
:func:`api.attachments.build_user_content`），文本对任何模型都可用，
不吃多模态能力。

**预算**：单份文档和单轮总量都有硬上限，超了截断并显式告诉模型"后面还有"，
免得一份 500 页 PDF 把 prompt 和账单一起撑爆。
"""

from __future__ import annotations

import codecs
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Final

from core.utils.logging import get_logger

logger = get_logger(__name__)

__all__ = [
    "ALLOWED_SUFFIXES",
    "DOC_SUFFIXES",
    "IMAGE_SUFFIXES",
    "KIND_DOCX",
    "KIND_IMAGE",
    "KIND_PDF",
    "KIND_TEXT",
    "KIND_XLSX",
    "MAX_DOC_BYTES",
    "MAX_EXTRACT_CHARS_PER_DOC",
    "MAX_EXTRACT_CHARS_TOTAL",
    "MIME_BY_SUFFIX",
    "DetectedFormat",
    "DocumentError",
    "ExtractedText",
    "detect_format",
    "document_block",
    "extract_text",
    "is_image_suffix",
    "mime_for_suffix",
]

# ── 类别 ──────────────────────────────────────────
KIND_IMAGE: Final = "image"
KIND_PDF: Final = "pdf"
KIND_DOCX: Final = "docx"
KIND_XLSX: Final = "xlsx"
KIND_TEXT: Final = "text"

#: 后缀 → MIME。以服务端**按魔数/扩展名判定的后缀**为准，不信客户端自报的
#: content_type。
MIME_BY_SUFFIX: Final[dict[str, str]] = {
    # 图片
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
    "gif": "image/gif",
    # 文档
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "txt": "text/plain",
    "md": "text/markdown",
    "csv": "text/csv",
    "json": "application/json",
    "log": "text/plain",
}

IMAGE_SUFFIXES: Final[frozenset[str]] = frozenset({"png", "jpg", "jpeg", "webp", "gif"})
DOC_SUFFIXES: Final[frozenset[str]] = frozenset(
    {"pdf", "docx", "xlsx", "txt", "md", "csv", "json", "log"}
)
ALLOWED_SUFFIXES: Final[frozenset[str]] = IMAGE_SUFFIXES | DOC_SUFFIXES

#: 期望是二进制容器的后缀（内容却可能是文本 —— 见 detect_format 里的兜底）
_OFFICE_SUFFIXES: Final[frozenset[str]] = frozenset({"pdf", "docx", "xlsx"})

#: 文档大小上限（图片仍是 10 MiB，见 api.storage.MAX_IMAGE_BYTES）
MAX_DOC_BYTES: Final[int] = 20 * 1024 * 1024

#: 单份文档抽出来的文本上限（字符）
MAX_EXTRACT_CHARS_PER_DOC: Final[int] = 40_000

#: 单轮对话里所有文档合计的文本上限（字符）
MAX_EXTRACT_CHARS_TOTAL: Final[int] = 80_000

#: xlsx 每张工作表最多取多少行
_MAX_SHEET_ROWS: Final[int] = 500
#: 每行最多取多少列（宽表只取前 N 列，避免一行几万个单元格）
_MAX_SHEET_COLS: Final[int] = 30

#: OOXML（docx/xlsx/pptx）都是 zip
_ZIP_MAGIC: Final = b"PK\x03\x04"
#: 旧版 Office 二进制格式（.doc/.xls/.ppt）
_CFB_MAGIC: Final = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class DocumentError(ValueError):
    """文档无法识别 / 解析失败。文案可直接展示给用户。"""


@dataclass(frozen=True, slots=True)
class DetectedFormat:
    """按魔数 + 扩展名判定的格式。``suffix`` 会被服务端用来拼存储 key。"""

    kind: str
    suffix: str
    mime: str


@dataclass(frozen=True, slots=True)
class ExtractedText:
    """一份文档抽出来的正文。"""

    kind: str
    text: str
    #: 是否因为超过字符预算被截断
    truncated: bool = False
    #: 人类可读的补充信息（页数 / 工作表名 / 段落数…）
    detail: str = ""


def mime_for_suffix(suffix: str) -> str:
    return MIME_BY_SUFFIX.get(suffix.lower(), "application/octet-stream")


def is_image_suffix(suffix: str) -> bool:
    return suffix.lower() in IMAGE_SUFFIXES


def _suffix_of(filename: str) -> str:
    return Path(filename or "").suffix.lstrip(".").lower()


def detect_format(head: bytes, filename: str) -> DetectedFormat:
    """按文件头魔数 + 扩展名判定格式。

    Args:
        head: 文件前若干字节（建议 ≥ 512）。
        filename: 客户端文件名（**只用来取扩展名**，不参与路径拼接）。

    Returns:
        :class:`DetectedFormat`。

    Raises:
        DocumentError: 不在白名单内，或魔数与扩展名对不上。
    """
    suffix = _suffix_of(filename)

    # 图片：完全靠魔数，扩展名随便写都能救回来
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return DetectedFormat(KIND_IMAGE, "png", MIME_BY_SUFFIX["png"])
    if head.startswith(b"\xff\xd8\xff"):
        return DetectedFormat(KIND_IMAGE, "jpg", MIME_BY_SUFFIX["jpg"])
    if head.startswith((b"GIF87a", b"GIF89a")):
        return DetectedFormat(KIND_IMAGE, "gif", MIME_BY_SUFFIX["gif"])
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return DetectedFormat(KIND_IMAGE, "webp", MIME_BY_SUFFIX["webp"])

    if head.startswith(b"%PDF-"):
        return DetectedFormat(KIND_PDF, "pdf", MIME_BY_SUFFIX["pdf"])

    if head.startswith(_CFB_MAGIC):
        # .doc/.xls/.ppt 是 OLE 复合文档，解析成本高、工具链也老，直接劝退
        raise DocumentError("不支持旧版 .doc / .xls / .ppt，请先另存为 .docx / .xlsx / .pdf")

    if head.startswith(_ZIP_MAGIC):
        # docx / xlsx / pptx 都是 zip，前几十字节里只有 [Content_Types].xml，
        # 分不出具体是哪种 —— 用扩展名初判，真正校验放在解析阶段（打不开就报错）。
        if suffix in {"docx", "xlsx"}:
            return DetectedFormat(suffix, suffix, MIME_BY_SUFFIX[suffix])
        raise DocumentError("这是 Office 压缩包，但扩展名不是 .docx / .xlsx")

    if suffix in {"txt", "md", "csv", "json", "log"}:
        return DetectedFormat(KIND_TEXT, suffix, MIME_BY_SUFFIX[suffix])

    # 扩展名写着 Office/PDF，内容其实是纯文本 —— 教务、ERP 系统的"导出 Excel"
    # 经常就是一坨 HTML 表格或 CSV 改了后缀。这种情况按文本收下，比甩一句
    # "不支持的文件类型"有用得多。
    if suffix in _OFFICE_SUFFIXES and _looks_like_text(head):
        logger.info("document.disguised_as_office", filename=filename, claimed=suffix)
        return DetectedFormat(KIND_TEXT, "txt", MIME_BY_SUFFIX["txt"])

    if suffix in _OFFICE_SUFFIXES:
        # 后缀像 Office，内容既不是 zip 也不是文本
        raise DocumentError(
            f"这个文件后缀是 .{suffix}，但内容不是真正的 {suffix.upper()} 文件"
            "（可能是别的格式改了后缀）。请用 Excel / Word 打开后另存为对应格式。"
        )

    allowed = "、".join(sorted(f".{s}" for s in ALLOWED_SUFFIXES))
    raise DocumentError(f"不支持的文件类型，仅支持：{allowed}")


def _looks_like_text(head: bytes) -> bool:
    """猜一段字节是不是纯文本。

    保守起见：先排除含 NUL 的（二进制特征），再要求能整段解码、且可打印字符
    占比够高。宁可漏判退回"不支持"，也别把二进制当文本喂给模型 —— 不危险，
    但模型会拿到一堆乱码。
    """
    if not head or b"\x00" in head:
        return False

    for encoding in ("utf-8", "gb18030"):
        try:
            text = head.decode(encoding)
        except UnicodeDecodeError:
            continue
        if not text:
            continue
        printable = sum(1 for ch in text if ch.isprintable() or ch in "\r\n\t")
        if printable >= len(text) * 0.95:
            return True
    return False


def extract_text(
    data: bytes,
    fmt: DetectedFormat,
    *,
    max_chars: int = MAX_EXTRACT_CHARS_PER_DOC,
) -> ExtractedText:
    """按格式抽正文。

    Args:
        data: 文件字节。
        fmt: :func:`detect_format` 的结果。
        max_chars: 抽取字符上限，超出截断。

    Returns:
        :class:`ExtractedText`。

    Raises:
        DocumentError: 文件损坏 / 加密 / 没有可提取的文字。
    """
    extractors = {
        KIND_PDF: _extract_pdf,
        KIND_DOCX: _extract_docx,
        KIND_XLSX: _extract_xlsx,
        KIND_TEXT: _extract_plain_text,
    }
    fn = extractors.get(fmt.kind)
    if fn is None:
        raise DocumentError(f"{fmt.suffix} 不是可提取文本的文档类型")
    return fn(data, max_chars=max_chars)


def document_block(extracted: ExtractedText, filename: str) -> str:
    """把抽取结果拼成给模型的文本块。"""
    header = f"[文档：{filename}" + (f"（{extracted.detail}）" if extracted.detail else "") + "]"
    body = extracted.text.strip() or "（没有可提取的文字内容）"
    tail = "\n（内容过长，已截断，后面还有内容）" if extracted.truncated else ""
    return f"{header}\n{body}{tail}"


# ── 各格式实现 ─────────────────────────────────────


def _finish(kind: str, text: str, max_chars: int, detail: str) -> ExtractedText:
    text = text.strip()
    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars]
    return ExtractedText(kind=kind, text=text, truncated=truncated, detail=detail)


def _extract_pdf(data: bytes, *, max_chars: int) -> ExtractedText:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # 空密码能开的（"仅限制编辑"）很常见，先试一下
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentError("PDF 有密码保护，请解密后再上传") from exc

        pages: list[str] = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
            except Exception:
                logger.warning("document.pdf.page_failed", page=i + 1)
                page_text = ""
            if page_text.strip():
                pages.append(page_text.strip())
        total = len(reader.pages)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"PDF 解析失败：{exc}") from exc

    if not pages:
        raise DocumentError("这个 PDF 里没有可提取的文字（可能是扫描件/纯图 PDF），请改传图片")

    return _finish(KIND_PDF, "\n\n".join(pages), max_chars, f"{total} 页")


def _extract_docx(data: bytes, *, max_chars: int) -> ExtractedText:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"Word 文档解析失败：{exc}") from exc

    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

    for index, table in enumerate(doc.tables, start=1):
        parts.append(f"\n[表格 {index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    detail = f"{len(doc.paragraphs)} 段" + (f" / {len(doc.tables)} 张表" if doc.tables else "")
    return _finish(KIND_DOCX, "\n".join(parts), max_chars, detail)


def _extract_xlsx(data: bytes, *, max_chars: int) -> ExtractedText:
    import openpyxl

    try:
        # read_only：按行流式读，别把大表整份读进内存
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError(f"Excel 解析失败：{exc}") from exc

    try:
        chunks: list[str] = []
        for sheet in workbook.worksheets:
            chunks.append(f"### 工作表：{sheet.title}")
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_no > _MAX_SHEET_ROWS:
                    chunks.append(f"…（本表只取前 {_MAX_SHEET_ROWS} 行）")
                    break
                cells = ["" if c is None else str(c) for c in row[:_MAX_SHEET_COLS]]
                line = "\t".join(cells).rstrip()
                if line.strip():
                    chunks.append(line)
            chunks.append("")
        names = "、".join(s.title for s in workbook.worksheets[:8])
    finally:
        workbook.close()

    detail = f"{len(workbook.worksheets)} 个工作表：{names}"
    return _finish(KIND_XLSX, "\n".join(chunks), max_chars, detail)


def _extract_plain_text(data: bytes, *, max_chars: int) -> ExtractedText:
    return _finish(KIND_TEXT, _decode_text(data), max_chars, "纯文本")


def _decode_text(data: bytes) -> str:
    """把文本文件解码成 str。

    顺序很重要：**UTF-8 → GB18030** 显式试，最后才交给 charset_normalizer。
    实测它对几百字节的短样本经常猜错（把 GBK 的 csv 猜成 EUC-KR），而
    Windows 上导出的 csv/txt 是 GBK 的概率很高。
    """
    if data.startswith(codecs.BOM_UTF8):
        return data.decode("utf-8-sig")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")

    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    from charset_normalizer import from_bytes

    best = from_bytes(data[:200_000]).best()
    if best is not None:
        return str(best)
    return data.decode("utf-8", errors="replace")
